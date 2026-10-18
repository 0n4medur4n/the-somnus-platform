"""Text out of uploaded originals (owner decision 2026-10-18).

The file is untrusted, so the tests are mostly about refusals: the type comes
from the bytes, and every limit refuses rather than half-reads.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from report.corpus import extraction
from report.corpus.extraction import (
    DOCX_CONTENT_TYPE,
    PDF_CONTENT_TYPE,
    ExtractionError,
    detect_content_type,
    extract,
)


def _pdf(*lines: str) -> bytes:
    """A real one-page PDF whose text layer holds `lines`."""
    text_ops = " ".join(
        f"72 {720 - 18 * index} Td ({line}) Tj 0 0 Td" if index == 0 else f"0 -18 Td ({line}) Tj"
        for index, line in enumerate(lines)
    )
    stream = f"BT /F1 12 Tf {text_ops} ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_pdf_is_read_into_plain_text() -> None:
    result = extract(_pdf("Higiene del sueno", "Rutina regular de descanso"))
    assert result.content_type == PDF_CONTENT_TYPE
    assert result.extension == "pdf"
    assert "Higiene del sueno" in result.text
    assert "Rutina regular de descanso" in result.text


def test_a_word_document_is_read_with_its_tables() -> None:
    data = _docx(
        ["Guía de higiene del sueño", "", "Mantén horarios regulares."],
        table=[["Hábito", "Frecuencia"], ["Ejercicio", "Diario"]],
    )
    result = extract(data)
    assert result.content_type == DOCX_CONTENT_TYPE
    assert result.extension == "docx"
    assert "Guía de higiene del sueño" in result.text
    assert "Mantén horarios regulares." in result.text
    assert "Ejercicio | Diario" in result.text


@pytest.mark.parametrize(
    "data",
    [
        b"\x89PNG\r\n\x1a\n" + b"\x00" * 64,
        b"MZ\x90\x00" + b"\x00" * 64,
        b"<html><body>hola</body></html>",
        b"plain text pretending to be a pdf",
    ],
)
def test_anything_but_pdf_or_docx_is_refused_by_its_bytes(data: bytes) -> None:
    with pytest.raises(ExtractionError):
        extract(data)


def test_a_zip_that_is_not_a_word_document_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hola")
    with pytest.raises(ExtractionError, match="only PDF and Word"):
        detect_content_type(buffer.getvalue())


def test_an_empty_file_is_refused() -> None:
    with pytest.raises(ExtractionError, match="empty"):
        extract(b"")


def test_a_file_over_the_size_limit_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    data = _pdf("Hola")
    monkeypatch.setattr(extraction, "MAX_FILE_BYTES", len(data) - 1)
    with pytest.raises(ExtractionError, match="larger"):
        extract(data)


def test_a_word_document_that_inflates_too_far_is_refused(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(extraction, "MAX_DOCX_UNCOMPRESSED_BYTES", 1024)
    with pytest.raises(ExtractionError, match="too large once opened"):
        extract(_docx(["Texto"]))


def test_a_password_protected_pdf_is_refused() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secreto")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ExtractionError, match="password"):
        extract(buffer.getvalue())


def test_a_pdf_with_no_text_layer_says_it_needs_ocr() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ExtractionError, match="OCR"):
        extract(buffer.getvalue())


def test_a_pdf_with_too_many_pages_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extraction, "MAX_PDF_PAGES", 0)
    with pytest.raises(ExtractionError, match="pages"):
        extract(_pdf("Hola"))


def test_a_broken_pdf_is_refused_not_crashed_on() -> None:
    with pytest.raises(ExtractionError):
        extract(b"%PDF-1.4\n" + b"\x00garbage" * 50)
