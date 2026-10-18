"""Plain text out of an uploaded PDF or Word file (owner decision 2026-10-18).

The extracted text becomes a DRAFT's text and then follows every corpus rule
unchanged: §B4's per-status cap at storage time, the rights gate at publish, and
embedding only on publish. Nothing here reaches a network or a model.

The file is untrusted input, so its type is decided by its bytes -- never by the
name or the declared content type -- and every parser runs under explicit
limits: pages, decompressed size, characters. A file that breaks a limit is
refused with a message the console can show, rather than half-read.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from typing import Final

#: The largest file the console may upload. Edge enforces the same number.
MAX_FILE_BYTES: Final = 20 * 1024 * 1024
#: A PDF longer than this is almost certainly not one reference document.
MAX_PDF_PAGES: Final = 600
#: A .docx is a zip; refuse one that would inflate past this (a zip bomb).
MAX_DOCX_UNCOMPRESSED_BYTES: Final = 100 * 1024 * 1024
#: Enough for a long guideline; the chunker splits it.
MAX_TEXT_CHARACTERS: Final = 1_000_000

PDF_CONTENT_TYPE: Final = "application/pdf"
DOCX_CONTENT_TYPE: Final = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CONTENT_TYPES: Final = (PDF_CONTENT_TYPE, DOCX_CONTENT_TYPE)


class ExtractionError(ValueError):
    """The file cannot become a document's text. The message is safe to show."""


@dataclass(frozen=True)
class ExtractedFile:
    content_type: str
    extension: str
    text: str


def detect_content_type(data: bytes) -> str:
    """PDF or DOCX by their bytes, or ExtractionError."""
    if data.startswith(b"%PDF-"):
        return PDF_CONTENT_TYPE
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return DOCX_CONTENT_TYPE
        except zipfile.BadZipFile as error:
            raise ExtractionError("the file is not a valid Word document") from error
    raise ExtractionError("only PDF and Word (.docx) files can be uploaded")


def _normalise(text: str) -> str:
    """Collapse the layout noise extraction leaves, keeping paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    # A hyphen at a line end is almost always a word broken across lines.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("the PDF is password-protected; upload an unprotected copy")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractionError(f"the PDF has more than {MAX_PDF_PAGES} pages")
        parts: list[str] = []
        total = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            total += len(page_text)
            if total > MAX_TEXT_CHARACTERS:
                raise ExtractionError("the document holds more text than one corpus entry can")
            parts.append(page_text)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as error:
        raise ExtractionError("the PDF could not be read") from error
    return "\n\n".join(parts)


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            inflated = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as error:
        raise ExtractionError("the file is not a valid Word document") from error
    if inflated > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise ExtractionError("the Word document is too large once opened")

    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except (ValueError, KeyError, OSError, zipfile.BadZipFile) as error:
        raise ExtractionError("the Word document could not be read") from error

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if len(text) > MAX_TEXT_CHARACTERS:
        raise ExtractionError("the document holds more text than one corpus entry can")
    return text


def extract(data: bytes) -> ExtractedFile:
    """The file's type and its plain text, or ExtractionError."""
    if not data:
        raise ExtractionError("the file is empty")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractionError(f"the file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    content_type = detect_content_type(data)
    if content_type == PDF_CONTENT_TYPE:
        raw, extension = _pdf_text(data), "pdf"
    else:
        raw, extension = _docx_text(data), "docx"
    text = _normalise(raw)
    if not text:
        raise ExtractionError(
            "no text could be read from the file; a scanned PDF needs OCR before it is uploaded"
        )
    return ExtractedFile(content_type=content_type, extension=extension, text=text)
